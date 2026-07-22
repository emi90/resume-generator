import io
from pathlib import Path

from docx import Document

from app.services.extractors.docx_extractor import extract_docx
from app.services.extractors.markdown_extractor import extract_markdown
from app.services.extractors.pdf_extractor import extract_pdf
from app.services.extractors.text_extractor import extract_text
from app.services.ingestion.resume_parser import parse_resume_file


FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures"


def _read_fixture_bytes(filename: str) -> bytes:
    return (FIXTURES_DIR / filename).read_bytes()



def test_text_extractor_utf8() -> None:
    result = extract_text(b"Senior Engineer\nPython")
    assert "Senior Engineer" in result.text



def test_markdown_extractor() -> None:
    result = extract_markdown(b"# Experience\n- Built APIs")
    assert "Experience" in result.text



def test_docx_extractor() -> None:
    document = Document()
    document.add_paragraph("Experience")
    document.add_paragraph("Built API services")

    stream = io.BytesIO()
    document.save(stream)

    result = extract_docx(stream.getvalue())
    assert "Built API services" in result.text


def test_docx_extractor_from_fixture() -> None:
    result = extract_docx(_read_fixture_bytes("sample_resume.docx"))
    assert isinstance(result.text, str)
    assert result.text.strip() != ""



def test_pdf_extractor_with_monkeypatch(monkeypatch) -> None:
    class FakePage:
        def __init__(self, text: str):
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, _):
            self.pages = [FakePage("A"), FakePage("B")]

    monkeypatch.setattr("app.services.extractors.pdf_extractor.PdfReader", FakeReader)

    result = extract_pdf(b"fake")
    assert result.text == "A\nB"
    assert result.meta["pages"] == 2


def test_pdf_extractor_from_fixture() -> None:
    result = extract_pdf(_read_fixture_bytes("sample_resume.pdf"))
    assert isinstance(result.text, str)
    assert result.meta["pages"] >= 1


def test_parse_resume_file_docx_fixture() -> None:
    normalized, warnings, file_type = parse_resume_file(
        "sample_resume.docx",
        _read_fixture_bytes("sample_resume.docx"),
    )
    assert file_type == "docx"
    assert normalized.strip() != ""
    assert isinstance(warnings, list)


def test_parse_resume_file_pdf_fixture() -> None:
    normalized, warnings, file_type = parse_resume_file(
        "sample_resume.pdf",
        _read_fixture_bytes("sample_resume.pdf"),
    )
    assert file_type == "pdf"
    assert isinstance(normalized, str)
    assert isinstance(warnings, list)


def test_parse_resume_file_xlsx_fixture_rejected() -> None:
    try:
        parse_resume_file(
            "sample_bad_resume.xlsx",
            _read_fixture_bytes("sample_bad_resume.xlsx"),
        )
    except Exception as exc:
        assert getattr(exc, "status_code", None) == 400
        assert getattr(exc, "code", None) == "UNSUPPORTED_RESUME_FILE_TYPE"
    else:
        raise AssertionError("Expected .xlsx fixture to be rejected as unsupported")



def test_parse_resume_file_rejects_unknown_type() -> None:
    try:
        parse_resume_file("resume.rtf", b"abc")
    except Exception as exc:
        assert getattr(exc, "status_code", None) == 400
    else:
        raise AssertionError("Expected unsupported file type to fail")
